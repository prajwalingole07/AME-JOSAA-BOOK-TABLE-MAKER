"""
Standalone one-institute cut-off Word builder.

Produces an A4 landscape .docx for a single IIT / NIT / IIIT using ONLY the
parsed Excel ranks (never the old Word file):

    - logo header (static/logo.jpg when available)
    - institute info box (state / address / phone / intake / fees ...)
    - rank table:  BRANCH (INTAKE) | GENDER | OPEN | OBC-NCL | EWS | SC | ST
                   one row per gender (Male / Female); the branch cell is
                   vertically merged so the name shows once, centred
    - quota divider rows (50% Home State / Other State) for NITs
    - missing ranks render as -----
"""

import io
import os
import re

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = os.path.join(BASE_DIR, "static", "logo.jpg")

CAT_GROUPS = [("OPEN", "OPEN-N", "OPEN-W"),
              ("OBC-NCL", "OBC-N", "OBC-W"),
              ("EWS", "EWS-N", "EWS-W"),
              ("SC", "SC-N", "SC-W"),
              ("ST", "ST-N", "ST-W")]

MISSING = "-----"
HEAD_BG = "1F4E79"
OPEN_BG = "2E75B6"
DIV_BG = "FFF2CC"
BRANCH_BG = "EDEDED"
MALE_BG = "DDEBF7"
FEMALE_BG = "FCE4EC"
MALE_RGB = (31, 105, 161)
FEMALE_RGB = (192, 0, 128)

_TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders",
               "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign",
               "hideMark"]


# --------------------------------------------------------------------------- #
#  tcPr helpers (insert elements in schema order)
# --------------------------------------------------------------------------- #
def _tcpr_add(cell, el):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = el.tag.split("}")[-1]
    pos = _TCPR_ORDER.index(tag) if tag in _TCPR_ORDER else len(_TCPR_ORDER)
    for child in tcPr:
        ctag = child.tag.split("}")[-1]
        cpos = _TCPR_ORDER.index(ctag) if ctag in _TCPR_ORDER else len(_TCPR_ORDER)
        if cpos > pos:
            child.addprevious(el)
            return
    tcPr.append(el)


def _shade(cell, color):
    _tcpr_add(cell, parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>'))


def _borders(cell, sz="4", color="808080"):
    _tcpr_add(cell, parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'))


def _margins(cell, top=20, bottom=20, left=40, right=40):
    _tcpr_add(cell, parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'))


def _vmerge(cell, restart=False):
    el = OxmlElement("w:vMerge")
    if restart:
        el.set(qn("w:val"), "restart")
    _tcpr_add(cell, el)


def _run(p, text, bold=False, size=8, color=None, font="Arial"):
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = font
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def _cell_text(cell, text, bold=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER,
               color=None, bg=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    _run(p, text, bold=bold, size=size, color=color)
    if bg:
        _shade(cell, bg)
    _borders(cell)
    _margins(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _cant_split(row):
    row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def _repeat_header(row):
    row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def safe_filename(name):
    s = re.sub(r'[\\/:*?"<>|]+', "", str(name or "")).strip()
    s = re.sub(r"\s+", " ", s)
    return (s.rstrip(". ") or "Institute") + ".docx"


def _rank_str(pair, idx):
    if not pair:
        return MISSING
    val = pair[idx]
    return MISSING if val is None else str(int(val))


# --------------------------------------------------------------------------- #
#  document parts
# --------------------------------------------------------------------------- #
def _init_document():
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width = Inches(11.69)
    sec.page_height = Inches(8.27)
    sec.top_margin = Inches(0.4)
    sec.bottom_margin = Inches(0.4)
    sec.left_margin = Inches(0.4)
    sec.right_margin = Inches(0.4)
    return doc


def _add_header(doc):
    if os.path.exists(LOGO_PATH):
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(LOGO_PATH, width=Inches(0.95))
        except Exception:
            pass
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    _run(p, "JOSAA 2026-27  |  OPENING & CLOSING RANKS  (Source: Excel)",
         bold=True, size=10.5, color=(31, 78, 121))


def _add_info_box(doc, inst):
    meta = inst.get("meta", {}) or {}
    name = meta.get("name_full") or inst.get("institute_name", "Institute")
    itype = meta.get("type", "")
    sr_no = inst.get("sr_no", 1)

    tbl = doc.add_table(rows=3, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    widths = [Inches(0.5), Inches(1.85), Inches(6.1), Inches(2.4)]

    c0 = tbl.cell(0, 0)
    c0.merge(tbl.cell(2, 0))
    _cell_text(c0, f"{sr_no}.", bold=True, size=13, bg="F2F2F2")

    _cell_text(tbl.cell(0, 1), "College Name & Type", bold=True, size=9,
               align=WD_ALIGN_PARAGRAPH.LEFT, bg="F2F2F2")
    _cell_text(tbl.cell(0, 2), f"{name}{f' ({itype})' if itype else ''}",
               bold=True, size=11, bg="F2F2F2")
    _cell_text(tbl.cell(0, 3), f"Established In {meta.get('established', '----')}",
               bold=True, size=9, align=WD_ALIGN_PARAGRAPH.LEFT, bg="F2F2F2")

    _cell_text(tbl.cell(1, 1), f"State: {meta.get('state', '----')}", size=9,
               align=WD_ALIGN_PARAGRAPH.LEFT, bg="F2F2F2")
    _cell_text(tbl.cell(1, 2), f"Address: {meta.get('address', '----')}", size=9,
               align=WD_ALIGN_PARAGRAPH.LEFT, bg="F2F2F2")
    _cell_text(tbl.cell(1, 3), f"Phone: {meta.get('phone', '----')}", size=9,
               align=WD_ALIGN_PARAGRAPH.LEFT, bg="F2F2F2")

    _cell_text(tbl.cell(2, 1), f"Total Intake: {meta.get('total_intake', '----')}",
               size=9, align=WD_ALIGN_PARAGRAPH.LEFT, bg="F2F2F2")
    _cell_text(tbl.cell(2, 2),
               f"Web: {meta.get('web', '----')}        "
               f"Email: {meta.get('email', '----')}",
               size=9, align=WD_ALIGN_PARAGRAPH.LEFT, bg="F2F2F2")
    _cell_text(tbl.cell(2, 3), f"Fees per Year: {meta.get('fees', '----')}",
               size=9, align=WD_ALIGN_PARAGRAPH.LEFT, bg="F2F2F2")

    for row in tbl.rows:
        _cant_split(row)
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = widths[i]

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(2)
    spacer.paragraph_format.space_after = Pt(2)


def _quota_sections(branches):
    order, groups = [], {}
    for b in branches:
        q = b.get("quota") or "AI"
        if q not in groups:
            groups[q] = []
            order.append(q)
        groups[q].append(b)
    return [(q, groups[q]) for q in order]


def _has_ranks(ranks, tokens):
    return any(ranks.get(t) for t in tokens)


def _add_rank_table(doc, branches):
    male_tokens = [n for _, n, _ in CAT_GROUPS]
    female_tokens = [w for _, _, w in CAT_GROUPS]

    sections = []
    for quota, rows in _quota_sections(branches):
        keep = [b for b in rows
                if _has_ranks(b.get("ranks") or {}, male_tokens + female_tokens)]
        if keep:
            sections.append((quota, keep))
    if not sections:
        return

    show_divider = any(q in ("HS", "OS") for q, _ in sections)

    body = []          # ("divider", label) | ("data", branch, [(gender, ranks, suffix)])
    for quota, rows in sections:
        if show_divider:
            body.append(("divider", {"HS": "50% Home State Quota",
                                     "OS": "50% Other State Quota"}.get(quota, "")))
        for b in rows:
            ranks = b.get("ranks") or {}
            pair = []
            if _has_ranks(ranks, male_tokens):
                pair.append(("Male", ranks, "N"))
            if _has_ranks(ranks, female_tokens):
                pair.append(("Female", ranks, "W"))
            if not pair:
                pair = [("Male", ranks, "N")]
            body.append(("data", b, pair))

    ncols = 2 + 2 * len(CAT_GROUPS)
    nrows = 2 + sum(1 if it[0] == "divider" else len(it[2]) for it in body)
    tbl = doc.add_table(rows=nrows, cols=ncols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    widths = [Inches(2.45), Inches(0.85)] + [Inches(0.755)] * 10

    # ---- header ------------------------------------------------------------
    hb = tbl.cell(0, 0).merge(tbl.cell(1, 0))
    _cell_text(hb, "BRANCH (INTAKE)", bold=True, size=8.5, bg=HEAD_BG,
               color=(255, 255, 255))
    gb = tbl.cell(0, 1).merge(tbl.cell(1, 1))
    _cell_text(gb, "GENDER", bold=True, size=8.5, bg=HEAD_BG,
               color=(255, 255, 255))
    for gi, (label, _, _) in enumerate(CAT_GROUPS):
        is_open = label == "OPEN"
        c1 = tbl.cell(0, 2 + gi * 2).merge(tbl.cell(0, 3 + gi * 2))
        _cell_text(c1, label, bold=True, size=10.5 if is_open else 9,
                   bg=OPEN_BG if is_open else HEAD_BG, color=(255, 255, 255))
        _cell_text(tbl.cell(1, 2 + gi * 2), "Open", bold=True, size=7.5,
                   bg=OPEN_BG if is_open else HEAD_BG, color=(255, 255, 255))
        _cell_text(tbl.cell(1, 3 + gi * 2), "Close", bold=True, size=7.5,
                   bg=OPEN_BG if is_open else HEAD_BG, color=(255, 255, 255))
    _repeat_header(tbl.rows[0])
    _repeat_header(tbl.rows[1])

    # ---- body --------------------------------------------------------------
    r = 2
    for item in body:
        if item[0] == "divider":
            row = tbl.rows[r]
            merged = row.cells[0].merge(row.cells[ncols - 1])
            _cell_text(merged, item[1], bold=True, size=9, bg=DIV_BG,
                       color=(128, 90, 0))
            _cant_split(row)
            r += 1
            continue

        _, branch, pair = item
        bcell = tbl.rows[r].cells[0]
        label = branch.get("display_name") or branch.get("program_name", "")
        _cell_text(bcell, label, bold=True, size=8,
                   align=WD_ALIGN_PARAGRAPH.LEFT, bg=BRANCH_BG)
        merged_branch = len(pair) > 1
        for idx, (gender, ranks, suffix) in enumerate(pair):
            row = tbl.rows[r]
            is_male = gender == "Male"
            _cell_text(row.cells[1], gender, bold=True, size=8,
                       color=MALE_RGB if is_male else FEMALE_RGB,
                       bg=MALE_BG if is_male else FEMALE_BG)
            for gi, (_, ntok, wtok) in enumerate(CAT_GROUPS):
                pair_r = ranks.get(ntok if suffix == "N" else wtok)
                for k in range(2):
                    _cell_text(row.cells[2 + gi * 2 + k], _rank_str(pair_r, k),
                               size=7.5)
            if merged_branch:
                if idx == 0:
                    _cell_text(row.cells[0], label, bold=True, size=8,
                               align=WD_ALIGN_PARAGRAPH.LEFT, bg=BRANCH_BG)
                    _vmerge(row.cells[0], restart=True)
                else:
                    _cell_text(row.cells[0], "", bg=BRANCH_BG)
                    _vmerge(row.cells[0], restart=False)
            _cant_split(row)
            r += 1

    for row in tbl.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = widths[i]


# --------------------------------------------------------------------------- #
#  public API
# --------------------------------------------------------------------------- #
def build_institute_docx(inst):
    """Return (bytes, filename) for one parser institute entry."""
    doc = _init_document()
    _add_header(doc)
    _add_info_box(doc, inst)
    _add_rank_table(doc, inst.get("branches") or [])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), safe_filename(inst.get("institute_name", "Institute"))
